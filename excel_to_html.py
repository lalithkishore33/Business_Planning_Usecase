from openpyxl import load_workbook
from html import escape

def excel_to_html(file_path, output_html="output.html"):
    wb = load_workbook(file_path, data_only=True)

    html = """
<!DOCTYPE html>
<html>
<head>
<meta charset="UTF-8">
<title>Excel Preview</title>

<style>

body{
    font-family:Arial;
    padding:20px;
}

table{
    border-collapse:collapse;
    margin-bottom:40px;
}

td{
    border:1px solid #999;
    padding:8px;
    min-width:100px;
    vertical-align:middle;
}

.sheet-title{
    font-size:22px;
    font-weight:bold;
    margin-bottom:15px;
}

</style>

</head>
<body>

"""

    for sheet in wb.worksheets:

        html += f'<div class="sheet-title">{escape(sheet.title)}</div>'
        html += "<table>"

        merged_ranges = {}

        # Store merge info
        for merged in sheet.merged_cells.ranges:
            min_col = merged.min_col
            min_row = merged.min_row
            max_col = merged.max_col
            max_row = merged.max_row

            merged_ranges[(min_row, min_col)] = (
                max_row - min_row + 1,
                max_col - min_col + 1
            )

        skip_cells = set()

        for merged in sheet.merged_cells.ranges:
            for r in range(merged.min_row, merged.max_row + 1):
                for c in range(merged.min_col, merged.max_col + 1):
                    if (r, c) != (merged.min_row, merged.min_col):
                        skip_cells.add((r, c))

        for row in sheet.iter_rows():

            html += "<tr>"

            for cell in row:

                if (cell.row, cell.column) in skip_cells:
                    continue

                rowspan = ""
                colspan = ""

                if (cell.row, cell.column) in merged_ranges:

                    rs, cs = merged_ranges[(cell.row, cell.column)]

                    if rs > 1:
                        rowspan = f' rowspan="{rs}"'

                    if cs > 1:
                        colspan = f' colspan="{cs}"'

                value = "" if cell.value is None else escape(str(cell.value))

                html += f"<td{rowspan}{colspan}>{value}</td>"

            html += "</tr>"

        html += "</table>"

    html += "</body></html>"

    with open(output_html, "w", encoding="utf-8") as f:
        f.write(html)

    print(f"HTML saved as {output_html}")


if __name__ == "__main__":
    excel_to_html("complex_financial_report.xlsx")
